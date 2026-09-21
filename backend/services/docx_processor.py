# DOCX Processor Engine using python-docx
# pyrefly: ignore [missing-import]
import docx
import os
from typing import Dict, Any, List

def process_docx_document(file_path: str) -> Dict[str, Any]:
    """
    Parses Word (.docx) documents into text paragraphs and tables.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found at {file_path}")

    doc = docx.Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    table_texts: List[str] = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                table_texts.append(row_text)

    full_text = "\n\n".join(paragraphs + table_texts)
    page_chunks = [{"page": 1, "text": full_text}]

    return {
        "page_count": 1,
        "pages": page_chunks,
        "full_text": full_text
    }
